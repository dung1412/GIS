import os
import io
import re
from typing import BinaryIO

try:
    import docx
except ImportError:
    pass # Will be handled by the user installing requirements

try:
    import openpyxl
except ImportError:
    pass

# TCVN3 to Unicode mapping
TCVN3_PAIRS = [
    # a
    (0xb5, 'à'), (0xb8, 'á'), (0xb6, 'ả'), (0xb7, 'ã'), (0xb9, 'ạ'),
    (0xa8, 'ă'), (0xbb, 'ằ'), (0xbe, 'ắ'), (0xbc, 'ẳ'), (0xbd, 'ẵ'), (0xc6, 'ặ'),
    (0xa9, 'â'), (0xc7, 'ầ'), (0xca, 'ấ'), (0xc8, 'ẩ'), (0xc9, 'ẫ'), (0xcb, 'ậ'),
    # e
    (0xcc, 'è'), (0xd0, 'é'), (0xce, 'ẻ'), (0xcf, 'ẽ'), (0xd1, 'ẹ'),
    (0xaa, 'ê'), (0xd2, 'ề'), (0xd5, 'ế'), (0xd3, 'ể'), (0xd4, 'ễ'), (0xd6, 'ệ'),
    # i
    (0xd7, 'ì'), (0xdd, 'í'), (0xd8, 'ỉ'), (0xdc, 'ĩ'), (0xde, 'ị'),
    # o
    (0xdf, 'ò'), (0xe3, 'ó'), (0xe1, 'ỏ'), (0xe2, 'õ'), (0xe4, 'ọ'),
    (0xab, 'ô'), (0xe5, 'ồ'), (0xe8, 'ố'), (0xe6, 'ổ'), (0xe7, 'ỗ'), (0xe9, 'ộ'),
    (0xac, 'ơ'), (0xea, 'ờ'), (0xed, 'ớ'), (0xeb, 'ở'), (0xec, 'ỡ'), (0xee, 'ợ'),
    # u
    (0xef, 'ù'), (0xf3, 'ú'), (0xf1, 'ủ'), (0xf2, 'ũ'), (0xf4, 'ụ'),
    (0xad, 'ư'), (0xf5, 'ừ'), (0xf8, 'ứ'), (0xf6, 'ử'), (0xf7, 'ữ'), (0xf9, 'ự'),
    # y
    (0xfa, 'ỳ'), (0xfd, 'ý'), (0xfb, 'ỷ'), (0xfc, 'ỹ'), (0xfe, 'ỵ'),
    # d
    (0xae, 'đ'), (0xa7, 'Đ'),
    # Base vowels uppercase
    (0xa1, 'Ă'), (0xa2, 'Â'), (0xa3, 'Ê'), (0xa4, 'Ô'), (0xa5, 'Ơ'), (0xa6, 'Ư')
]

TCVN3_UPPER_SEQUENCES = [
    ('A\xb5', 'À'), ('A\xb8', 'Á'), ('A\xb6', 'Ả'), ('A\xb7', 'Ã'), ('A\xb9', 'Ạ'),
    ('\xa1\xbb', 'Ằ'), ('\xa1\xbe', 'Ắ'), ('\xa1\xbc', 'Ẳ'), ('\xa1\xbd', 'Ẵ'), ('\xa1\xc6', 'Ặ'),
    ('\xa2\xc7', 'Ầ'), ('\xa2\xca', 'Ấ'), ('\xa2\xc8', 'Ẩ'), ('\xa2\xc9', 'Ẫ'), ('\xa2\xcb', 'Ậ'),
    ('E\xcc', 'È'), ('E\xd0', 'É'), ('E\xce', 'Ẻ'), ('E\xcf', 'Ẽ'), ('E\xd1', 'Ẹ'),
    ('\xa3\xd2', 'Ề'), ('\xa3\xd5', 'Ế'), ('\xa3\xd3', 'Ể'), ('\xa3\xd4', 'Ễ'), ('\xa3\xd6', 'Ệ'),
    ('I\xd7', 'Ì'), ('I\xdd', 'Í'), ('I\xd8', 'Ỉ'), ('I\xdc', 'Ĩ'), ('I\xde', 'Ị'),
    ('O\xdf', 'Ò'), ('O\xe3', 'Ó'), ('O\xe1', 'Ỏ'), ('O\xe2', 'Õ'), ('O\xe4', 'Ọ'),
    ('\xa4\xe5', 'Ồ'), ('\xa4\xe8', 'Ố'), ('\xa4\xe6', 'Ổ'), ('\xa4\xe7', 'Ỗ'), ('\xa4\xe9', 'Ộ'),
    ('\xa5\xea', 'Ờ'), ('\xa5\xed', 'Ớ'), ('\xa5\xeb', 'Ở'), ('\xa5\xec', 'Ỡ'), ('\xa5\xee', 'Ợ'),
    ('U\xef', 'Ù'), ('U\xf3', 'Ú'), ('U\xf1', 'Ủ'), ('U\xf2', 'Ũ'), ('U\xf4', 'Ụ'),
    ('\xa6\xf5', 'Ừ'), ('\xa6\xf8', 'Ứ'), ('\xa6\xf6', 'Ử'), ('\xa6\xf7', 'Ữ'), ('\xa6\xf9', 'Ự'),
    ('Y\xfa', 'Ỳ'), ('Y\xfd', 'Ý'), ('Y\xfb', 'Ỷ'), ('Y\xfc', 'Ỹ'), ('Y\xfe', 'Ỵ'),
    ('\xad\xac', 'ươ'), ('\xad\xed', 'ướ'), ('\xad\xea', 'ườ'), ('\xad\xeb', 'ưở'), ('\xad\xec', 'ưỡ'), ('\xad\xee', 'ượ')
]

tcvn3_to_uni_map = {chr(code): uni for code, uni in TCVN3_PAIRS}
uni_to_tcvn3_map = {uni: chr(code) for code, uni in TCVN3_PAIRS}
for tcvn, uni in TCVN3_UPPER_SEQUENCES:
    uni_to_tcvn3_map[uni] = tcvn

TCVN3_HALLMARKS = {
    0xa1, 0xa2, 0xa3, 0xa4, 0xa5, 0xa6, 0xa7, 0xa8, 0xa9, 0xaa, 0xab, 0xac, 0xad, 0xae,
    0xb5, 0xb6, 0xb7, 0xb8, 0xb9, 0xbb, 0xbc, 0xbd, 0xbe,
    0xc6, 0xd7, 0xd8, 0xdc, 0xde, 0xdf, 0xf7, 0xfe
}

def is_already_unicode(text: str) -> bool:
    for char in text:
        cp = ord(char)
        if cp in (0x0110, 0x0111, 0x01a0, 0x01a1, 0x01af, 0x01b0) or (0x1ea0 <= cp <= 0x1ef9):
            return True
    return False

def is_tcvn3(text: str) -> bool:
    if is_already_unicode(text):
        return False
    for char in text:
        if ord(char) in TCVN3_HALLMARKS:
            return True
    for seq, _ in TCVN3_UPPER_SEQUENCES:
        if seq in text:
            return True
    return False

def tcvn3_to_unicode(text: str) -> str:
    if not text or not isinstance(text, str):
        return text or ""
    if not is_tcvn3(text):
        import unicodedata
        return unicodedata.normalize('NFC', text)
    
    tokens = re.split(r'(\s+|[^\w\u00a1-\u00fe]+)', text)
    out = []
    for tok in tokens:
        if not tok:
            continue
        ascii_letters = [c for c in tok if c.isalpha() and ord(c) < 128]
        is_all_caps = len(ascii_letters) >= 2 and all(c.isupper() for c in ascii_letters)
        
        res = tok
        for seq, uni in TCVN3_UPPER_SEQUENCES:
            res = res.replace(seq, uni)
        
        word_out = "".join(tcvn3_to_uni_map.get(ch, ch) for ch in res)
        out.append(word_out.upper() if is_all_caps else word_out)
        
    import unicodedata
    return unicodedata.normalize('NFC', "".join(out))

def unicode_to_tcvn3(text: str) -> str:
    if not text or not isinstance(text, str):
        return text or ""
    out = "".join(uni_to_tcvn3_map.get(ch, ch) for ch in text)
    return out

def convert_text(text: str, to_unicode: bool) -> str:
    return tcvn3_to_unicode(text) if to_unicode else unicode_to_tcvn3(text)

def convert_docx(file_bytes: bytes, to_unicode: bool = True) -> bytes:
    import docx
    doc = docx.Document(io.BytesIO(file_bytes))
    for para in doc.paragraphs:
        for run in para.runs:
            if run.text:
                run.text = convert_text(run.text, to_unicode)
                if run.font.name and ('VNTime' in run.font.name or 'ABC' in run.font.name or '.Vn' in run.font.name):
                    run.font.name = 'Times New Roman' if to_unicode else '.VnTime'
    
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        if run.text:
                            run.text = convert_text(run.text, to_unicode)
                            if run.font.name and ('VNTime' in run.font.name or 'ABC' in run.font.name or '.Vn' in run.font.name):
                                run.font.name = 'Times New Roman' if to_unicode else '.VnTime'
                                
    out_stream = io.BytesIO()
    doc.save(out_stream)
    return out_stream.getvalue()

def convert_xlsx(file_bytes: bytes, to_unicode: bool = True) -> bytes:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))
    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    cell.value = convert_text(cell.value, to_unicode)
                    if cell.font and cell.font.name and ('.Vn' in cell.font.name or 'VNTime' in cell.font.name):
                        cell.font = openpyxl.styles.Font(name='Times New Roman' if to_unicode else '.VnTime', 
                                                         size=cell.font.size, 
                                                         bold=cell.font.bold, 
                                                         italic=cell.font.italic)
    
    out_stream = io.BytesIO()
    wb.save(out_stream)
    return out_stream.getvalue()
